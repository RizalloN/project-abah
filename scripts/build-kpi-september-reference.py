"""Build read-only KPI snapshots from the supplied September workbooks.

Usage: python scripts/build-kpi-september-reference.py <source-directory>
Run MarkItDown inspection before using this builder. Cached Excel results are
retained; formulas are never recalculated and no source workbook is written.
"""
import hashlib
import json
import subprocess
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    "mantri": ("KPI Mantri Tumbuh 30 September 2026.xlsx", "KPI Mantri Tumbuh", 1, 2, 1, 7),
    "mantri-briguna": ("KPI Mantri Briguna 30 September 2026.xlsx", "KPI Mantri Briguna", 1, 2, 1, 7),
    "ka-unit": ("KPI KAUNIT AREA 6 Madiun update 30 september 2026.xlsx", "KPI Kaunit", 3, 4, 0, 3),
    "mbm": ("KPI MBM AREA 6 Madiun update 30 september 2026.xlsx", "KPI MBM", 3, 4, 0, 1),
}
FORMATTER = r'''
require 'vendor/autoload.php';
$input = json_decode(stream_get_contents(STDIN), true, 512, JSON_THROW_ON_ERROR);
$rows = [];
foreach ($input['values'] as $r => $row) {
    $rows[] = array_map(static function ($value, $format) {
        if ($value === null) return '';
        if (!is_numeric($value) || is_string($value)) return trim((string) $value);
        return trim(\PhpOffice\PhpSpreadsheet\Style\NumberFormat::toFormattedString($value, $format));
    }, $row, $input['formats'][$r]);
}
echo json_encode($rows, JSON_THROW_ON_ERROR | JSON_UNESCAPED_UNICODE);
'''


def build(directory):
    reports = []
    for key, (filename, sheet_name, header_row, sub_row, bo_index, identity_index) in SOURCES.items():
        source = directory / filename
        book = load_workbook(source, read_only=True, data_only=True)
        sheet = book[sheet_name]
        sheet.reset_dimensions()
        source_rows = {}
        for row_number, cells in enumerate(sheet.iter_rows(), 1):
            if any(c.value is not None for c in cells):
                source_rows[row_number] = [(c.value, c.number_format, c.data_type) for c in cells]
        width = max(len(cells) for cells in source_rows.values())
        padded = lambda row: list(row) + [(None, 'General', 'n')] * (width - len(row))
        headers = [[str(c[0] or '').strip() for c in padded(source_rows[n])] for n in [header_row, sub_row]]
        if key in ['mbm', 'ka-unit']:
            for i, weight in enumerate(headers[1]):
                if weight.startswith('BOBOT '):
                    headers[0][i] += ' ' + weight
                    headers[1][i] = 'Pencapaian'
                    headers[1][i + 1] = 'Score'
        row_numbers, values, formats, branches, errors = [], [], [], Counter(), []
        for row_number, cells in source_rows.items():
            if row_number <= sub_row:
                continue
            cells = padded(cells)
            branch = str(cells[bo_index][0] or '').strip()
            if not branch:
                continue  # Blank control/formula rows have no personnel identity.
            if branch not in ['MADIUN', 'MAGETAN', 'NGAWI', 'PONOROGO'] or not cells[identity_index][0]:
                raise ValueError(f'Unrecognised personnel row {sheet_name}!{row_number}: {branch}')
            row_numbers.append(row_number)
            values.append([c[0] for c in cells])
            formats.append([c[1] for c in cells])
            branches[branch] += 1
            for i, cell in enumerate(cells):
                if cell[2] == 'e':
                    errors.append({'cell': f'{get_column_letter(i + 1)}{row_number}', 'value': cell[0]})
        formatted = subprocess.run(['php', '-r', FORMATTER], cwd=ROOT, input=json.dumps({'values': values, 'formats': formats}), text=True, encoding='utf-8', capture_output=True, check=True)
        rows = json.loads(formatted.stdout)
        format_list = list(dict.fromkeys(f for row in formats for f in row))
        snapshot = {
            'version': 1, 'period': '2026-09', 'as_of': '2026-09-30',
            'source': {'file': filename, 'sha256': hashlib.file_digest(source.open('rb'), 'sha256').hexdigest(), 'sheet': sheet_name},
            'header_rows': headers, 'source_columns': [get_column_letter(i + 1) for i in range(width)],
            'source_rows': row_numbers, 'rows': rows, 'raw_rows': values,
            'number_formats': format_list, 'format_ids': [[format_list.index(f) for f in row] for row in formats],
            'branches': dict(branches), 'source_errors': errors,
        }
        destination = ROOT / 'resources/data' / f'kpi-{key}-2026-09.json'
        destination.write_text(json.dumps(snapshot, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
        reports.append({'sheet': key, 'rows': len(rows), 'columns': width, 'branches': dict(branches), 'errors': len(errors)})
        book.close()
    print(json.dumps(reports))


if __name__ == '__main__':
    build(Path(sys.argv[1]))
